"""Regression tests for CV parsing in anvil/ingest.py: role headers, DOCX list and
table order, PDF wrapped lines, and the headline position."""
from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from anvil import ingest  # noqa: E402

SKELETON = json.loads((ROOT / "vault" / "skeleton.json").read_text(encoding="utf-8"))
ROLE_EY = "Strategy Consultant – EY – Apr 2025 – May 2025"


def _numbered(paragraph) -> None:
    """Give a paragraph direct numbering (w:numPr in its pPr)."""
    pPr = paragraph._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numId = OxmlElement("w:numId")
    numId.set(qn("w:val"), "1")
    numPr.append(ilvl)
    numPr.append(numId)
    pPr.append(numPr)


class RoleHeaderTests(unittest.TestCase):
    def test_company_line_without_pipe_is_a_role_header(self):
        lines = ["SWASTIK KUMAR RAY", "Product Analyst", "Experience", ROLE_EY,
                 "Built the forecast model used by the client team"]
        cv = ingest.parse_cv(lines, SKELETON)
        self.assertEqual(cv["unknown_roles"], [])
        self.assertEqual(cv["experience"], {"ey": ["Built the forecast model used by the client team"]})

    def test_bullet_mentioning_company_is_not_a_role_header(self):
        lines = ["NAME", "Headline", "Experience", ROLE_EY, "- Supported EY partners on a pricing study"]
        cv = ingest.parse_cv(lines, SKELETON)
        self.assertEqual(cv["experience"]["ey"], ["Supported EY partners on a pricing study"])


class HeadlineTests(unittest.TestCase):
    def test_first_non_contact_line_is_headline(self):
        lines = ["SWASTIK KUMAR RAY", "Odisha, India | +91 63717 56663 | a@b.com",
                 "Product Analyst | CTRM", "Nagpur, India", "Experience"]
        cv = ingest.parse_cv(lines, SKELETON)
        self.assertEqual(cv["headline"], "Product Analyst | CTRM")


class DocxTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def _save(self, doc, name="cv.docx") -> Path:
        path = self.dir / name
        doc.save(str(path))
        return path

    def test_list_paragraph_heading_is_not_a_bullet(self):
        doc = Document()
        doc.add_paragraph("SWASTIK KUMAR RAY")
        doc.add_paragraph("Key Strengths", style="List Paragraph")
        doc.add_paragraph("Plain sentence about the work.")
        lines = ingest.read_lines(self._save(doc), SKELETON)
        self.assertEqual(lines[1], "Key Strengths")

    def test_numbered_and_bullet_glyph_paragraphs_are_bullets(self):
        doc = Document()
        doc.add_paragraph("SWASTIK KUMAR RAY")
        p = doc.add_paragraph("Numbered achievement")
        _numbered(p)
        doc.add_paragraph("• Glyph achievement")
        lines = ingest.read_lines(self._save(doc), SKELETON)
        self.assertEqual(lines[1], "- Numbered achievement")
        self.assertEqual(lines[2], "• Glyph achievement")

    def test_role_header_in_list_style_is_not_a_bullet(self):
        doc = Document()
        doc.add_paragraph("SWASTIK KUMAR RAY")
        doc.add_paragraph("Experience")
        doc.add_paragraph(ROLE_EY, style="List Paragraph")
        doc.add_paragraph("Built the forecast model")
        path = self._save(doc)
        lines = ingest.read_lines(path, SKELETON)
        self.assertEqual(lines[2], ROLE_EY)
        cv = ingest.parse_cv(lines, SKELETON)
        self.assertEqual(cv["experience"]["ey"], ["Built the forecast model"])

    def test_table_lines_keep_document_order(self):
        doc = Document()
        doc.add_paragraph("SWASTIK KUMAR RAY")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Product Analyst"
        table.rows[0].cells[1].text = "CTRM"
        doc.add_paragraph("Summary")
        doc.add_paragraph("Closing paragraph")
        lines = ingest.read_lines(self._save(doc), SKELETON)
        self.assertEqual(lines, ["SWASTIK KUMAR RAY", "Product Analyst", "CTRM", "Summary", "Closing paragraph"])
        cv = ingest.parse_cv(lines, SKELETON)
        self.assertEqual(cv["headline"], "Product Analyst")


class PdfJoinTests(unittest.TestCase):
    def join(self, raw):
        return ingest._join_wrapped(raw, SKELETON)

    def test_wrapped_bullet_is_joined_and_whitespace_collapsed(self):
        out = self.join(["Experience", "• Built a pricing model that cut   manual",
                         "work  by half for the team", "• Second achievement."])
        self.assertEqual(out, ["Experience", "• Built a pricing model that cut manual work by half for the team",
                               "• Second achievement."])

    def test_bullet_ending_with_punctuation_is_not_joined(self):
        out = self.join(["Experience", "• First sentence.", "Standalone continuation"])
        self.assertEqual(out, ["Experience", "• First sentence.", "Standalone continuation"])

    def test_new_bullet_and_role_header_are_not_joined(self):
        out = self.join(["Experience", "• Did a thing", "• Did another thing", ROLE_EY, "• Work"])
        self.assertEqual(out, ["Experience", "• Did a thing", "• Did another thing", ROLE_EY, "• Work"])

    def test_skill_group_continuation_joins_previous_group(self):
        out = self.join(["Skills", "Languages: Python, SQL,", "R and Julia", "Tools: Power BI"])
        self.assertEqual(out, ["Skills", "Languages: Python, SQL, R and Julia", "Tools: Power BI"])

    def test_pdf_lines_end_to_end_parse(self):
        raw = ["SWASTIK KUMAR RAY", "Product Analyst", "Experience", ROLE_EY,
               "• Built the forecast model that the client", "team now uses every week"]
        cv = ingest.parse_cv(ingest._join_wrapped(raw, SKELETON), SKELETON)
        self.assertEqual(cv["experience"]["ey"], ["Built the forecast model that the client team now uses every week"])


if __name__ == "__main__":
    unittest.main()
