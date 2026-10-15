"""Run: python -m unittest discover tests"""
import json
import os
import random
import shutil
import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parents[1]


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        shutil.copytree(ROOT / "vault", self.tmp / "vault",
                        ignore=shutil.ignore_patterns("gene-pool.json", "overrides.json"))
        os.environ["ANVIL_HOME"] = str(self.tmp)
        from anvil import ingest, store
        self.store = store
        self.pool = ingest.run()
        self.sk = store.skeleton()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)
        os.environ.pop("ANVIL_HOME", None)


class TestGuards(unittest.TestCase):
    def setUp(self):
        self.sk = json.loads((ROOT / "vault" / "skeleton.json").read_text())

    def test_banned_words_and_em_dash(self):
        from anvil import guards
        r = guards.check_style("Spearheaded a cutting-edge rollout — fast", self.sk)
        rules = {v.detail for v in r.violations}
        self.assertTrue(any("spearheaded" in d for d in rules))
        self.assertTrue(any("cutting-edge" in d for d in rules))
        self.assertTrue(any("—" in d for d in rules))
        self.assertTrue(guards.check_style("Cleared 7 of 17 defects – en dash is fine", self.sk).ok)

    def test_no_new_numbers_or_nouns(self):
        from anvil import guards
        src = "Cleared 7 of 17 queued defects in a 9-hour release-crunch session"
        self.assertTrue(guards.assert_no_new_facts(src, "Resolved 7 of 17 queued defects in a 9-hour release crunch").ok)
        self.assertFalse(guards.assert_no_new_facts(src, "Cleared 8 of 17 queued defects").ok)
        self.assertFalse(guards.assert_no_new_facts(src, "Cleared 7 of 17 Salesforce defects").ok)
        self.assertTrue(guards.assert_no_new_facts(src, "Cleared 7 of 17 Acme defects", allowed_extra="Acme").ok)

    def test_petronas_scope(self):
        from anvil import guards
        self.assertTrue(guards.check_client_scope("Petronas – trade data import/export", self.sk).ok)
        self.assertFalse(guards.check_client_scope("Configured pricing for Petronas", self.sk).ok)
        # per clause: CBH testing and Petronas import/export in one bullet is fine
        self.assertTrue(guards.check_client_scope(
            "CBH – testing, pricing checks; Petronas – trade data import/export", self.sk).ok)


class TestIngest(Base):
    def test_pool_shape(self):
        kinds = {}
        for g in self.pool["genes"].values():
            kinds[g["kind"]] = kinds.get(g["kind"], 0) + 1
        self.assertEqual(kinds["bullet"], 15)  # 7 Fendahl + 3 EY + 5 Startup Odisha, deduped across 4 CVs
        self.assertEqual(kinds["summary"], 4)
        self.assertEqual(kinds["project"], 2)
        blocked = [g for g in self.pool["genes"].values() if g["blocked"]]
        self.assertEqual(len(blocked), 1)  # commodity summary links Petronas to pricing
        self.assertIn("Petronas", blocked[0]["block_reasons"][0])

    def test_docx_roundtrip(self):
        import docx
        from anvil import ingest
        d = docx.Document()
        for line in ["SWASTIK KUMAR RAY", "Odisha, India | +91 63717 56663 | swastikray90@gmail.com",
                     "PROFESSIONAL SUMMARY", "Product Analyst with a new real summary line.",
                     "EXPERIENCE", "Product Analyst – Fendahl (Fusion CTRM), Nagpur | Apr 2026 – Present"]:
            d.add_paragraph(line)
        d.add_paragraph("Built product demo environments and materials for prospective clients in pre-sales cycles.",
                        style="List Bullet")
        p = self.tmp / "vault" / "sources" / "05-test.docx"
        d.save(p)
        pool = ingest.run()
        texts = [g["text"] for g in pool["genes"].values()]
        self.assertIn("Product Analyst with a new real summary line.", texts)
        demo = [g for g in pool["genes"].values() if g["text"].startswith("Built product demo")]
        self.assertEqual(len(demo), 1)
        self.assertIn("05-test.docx", [s["file"] for s in demo[0]["sources"]])


class TestTailorAndBandit(Base):
    JD = ("Business Analyst, Implementation. You will gather requirements, run UAT, manage defects and "
          "support data migration for commodity trading clients. SQL, Excel, Tableau, SAP. Stakeholder "
          "management. We help firms streamline onboarding and improve data quality. " * 3)

    def args(self, **kw):
        jd = self.tmp / "jd.txt"
        jd.write_text(self.JD)
        base = dict(jd_file=str(jd), company="Acme Trade", role="Business Analyst", url="", leads_file=None,
                    company_url=None, notes=None, cluster=None, polish=False, golden=False,
                    formats="html", seed=3, refresh=False, dry_run=False, force=False, copy_to=None)
        base.update(kw)
        return SimpleNamespace(**base)

    def _no_ember(self):
        from anvil import ember
        ember.company_url_candidates = lambda lead, jd="": []

    def test_tailor_writes_guarded_cv_and_logs(self):
        self._no_ember()
        from anvil import bandit, tailor
        app = tailor.run("manual:acme:1", self.args())["application"]
        self.assertEqual(app["outcome"], "PENDING")
        self.assertEqual(app["cluster"], "ba_pa")
        out = self.tmp / app["out_dir"]
        html = (out / "Swastik_Kumar_Ray_Acme_Trade.html").read_text()
        self.assertIn("Nagpur, Maharashtra, India", html)
        self.assertIn("Product Analyst", html)
        self.assertNotIn("—", html)
        self.assertNotIn("Tableau", html)  # true gap never added
        diff = (out / "jd-diff.md").read_text()
        self.assertIn("Tableau", diff.split("## 3.")[1])
        recon = json.loads((out / "manual_acme_1-recon.json").read_text())
        self.assertEqual(recon["dna_hash"], app["dna_hash"])
        self.assertEqual(len(bandit.load_apps()), 1)

    def test_feedback_moves_weights_and_golden(self):
        self._no_ember()
        from anvil import bandit, tailor
        app = tailor.run("manual:acme:1", self.args())["application"]
        bandit.apply_outcome("manual:acme:1", "callback")
        log = bandit.load_log()
        g = app["genes_used"][0]
        self.assertEqual(log["genes"]["ba_pa"][g]["alpha"], 3)
        self.assertEqual(log["golden"]["ba_pa"][0]["dna_hash"], app["dna_hash"])

        app2 = tailor.run("manual:acme:2", self.args(company="Beta Co"))["application"]
        bandit.apply_outcome("manual:acme:2", "rejected")
        log = bandit.load_log()
        for gid in app2["genes_used"]:
            self.assertGreaterEqual(log["genes"]["ba_pa"][gid]["beta"], 2)

    def test_ghosting_and_deprecation(self):
        self._no_ember()
        from anvil import bandit, engine, tailor
        apps = []
        for i in range(8):
            apps.append(tailor.run(f"manual:co{i}:1", self.args(company=f"Co{i}", seed=1))["application"])
        old = (datetime.now(timezone.utc) - timedelta(days=20)).isoformat(timespec="seconds")
        data = bandit.load_apps()
        for a in data:
            a["applied_at"] = old
        bandit.save_apps(data)
        ghosted = bandit.ghost_sweep()
        self.assertEqual(len(ghosted), 8)
        log = bandit.load_log()
        hooks = log["hooks"]["ba_pa"]
        dep = [h for h, v in hooks.items() if v["deprecated"]]
        self.assertTrue(dep, "a hook sent 8 times with 0 replies must be deprecated")
        nxt = tailor.run("manual:co9:1", self.args(company="Co9", seed=1))["application"]
        self.assertNotIn(nxt["hook_id"], dep)

    def test_polish_rejects_invented_facts(self):
        self._no_ember()
        from anvil import engine, llm
        pool, sk = self.pool, self.sk
        rel = engine.relevance(pool, self.JD)
        from anvil import bandit
        plan = engine.choose(pool, sk, "ba_pa", rel, bandit.load_log(), random.Random(1))
        before = dict(plan.texts)
        llm.polish = lambda lines, t, a: [ln + " using Salesforce for 40 clients" for ln in lines]
        engine.apply_polish(plan, pool, sk, {}, [], "Acme", "BA")
        self.assertEqual(plan.texts, before)
        self.assertTrue(any("rejected" in c for c in plan.changes))


class TestEmberLead(Base):
    def test_lead_from_csv_with_bot_walled_company_site(self):
        from anvil import ember, recon, tailor
        fx = ROOT / "tests" / "fixtures"

        def fake(url, max_tokens=2500, timeout=60):
            if "lever.co" in url:
                return {"title": "job", "markdown": (fx / "sprinto-jd.md").read_text()}
            if "sprinto.com" in url:
                return {"title": "Just a moment...", "markdown": "Checking your browser"}
            raise ember.EmberError("404")
        ember.scrape = fake
        cid = "lever:sprinto:56d536c2-812c-49be-8421-6ad63c42a273"
        args = SimpleNamespace(jd_file=None, company=None, role=None, url=None, leads_file=str(fx / "open-leads.csv"),
                               company_url=None, notes="Hiring manager posted about fixing HubSpot hygiene.",
                               cluster=None, polish=False, golden=False, formats="html", seed=2, refresh=False,
                               dry_run=False, force=False, copy_to=None)
        app = tailor.run(cid, args)["application"]
        self.assertEqual(app["cluster"], "revops_ops")
        rec = recon.load_recon(cid)
        self.assertTrue(rec["product_summary"][0].startswith("Sprinto is"))
        self.assertIn("HubSpot", [g["skill"] for g in rec["jd_diff"]["gaps"]])
        self.assertIn("HubSpot hygiene", rec["notes"])
        # DEAD leads are refused
        with self.assertRaises(SystemExit):
            tailor.run("lever:oldco:1", args)["application"]


class TestGitHubState(unittest.TestCase):
    def test_branch_created_and_sha_conflict_retried(self):
        import base64
        from anvil import store
        files, branches, calls = {}, {"main"}, []

        def fake(self, method, path, body=None):
            calls.append((method, path))
            if path.startswith("contents/state/"):
                rel = path.split("contents/state/")[1].split("?")[0]
                if method == "GET":
                    if "anvil-state" not in branches or rel not in files:
                        return 404, {}
                    sha, data = files[rel]
                    return 200, {"sha": sha, "content": base64.b64encode(data.encode()).decode()}
                if "anvil-state" not in branches:
                    return 404, {}
                if rel in files and body.get("sha") != files[rel][0]:
                    return 409, {}
                sha = f"s{len(calls)}"
                files[rel] = (sha, base64.b64decode(body["content"]).decode())
                return 200, {"content": {"sha": sha}}
            if path == "branches/anvil-state":
                return (200 if "anvil-state" in branches else 404), {}
            if path == "":
                return 200, {"default_branch": "main"}
            if path == "git/ref/heads/main":
                return 200, {"object": {"sha": "abc"}}
            if path == "git/refs":
                branches.add("anvil-state")
                return 201, {}
            return 404, {}

        os.environ.update(ANVIL_STATE="github", GITHUB_TOKEN="t")
        self.addCleanup(lambda: [os.environ.pop(k, None) for k in ("ANVIL_STATE", "GITHUB_TOKEN")])
        orig = store._GitHub._req
        store._GitHub._req = fake
        self.addCleanup(setattr, store._GitHub, "_req", orig)
        store._gh = None
        self.assertIsNone(store.read_state("applications.json"))
        store.write_state("applications.json", [{"a": 1}])
        self.assertIn("anvil-state", branches)
        self.assertEqual(store.read_state("applications.json"), [{"a": 1}])
        files["applications.json"] = ("other", "[]")  # someone else wrote in between
        store.write_state("applications.json", [{"a": 2}])
        self.assertEqual(store.read_state("applications.json"), [{"a": 2}])
        store._gh = None


class TestServer(Base):
    def call(self, method, path, body=None, key="k1", qs=""):
        import io
        from anvil import server
        raw = json.dumps(body).encode() if body is not None else b""
        env = {"REQUEST_METHOD": method, "PATH_INFO": path, "QUERY_STRING": qs, "CONTENT_LENGTH": str(len(raw)),
               "wsgi.input": io.BytesIO(raw), "HTTP_HOST": "x.test", "wsgi.url_scheme": "https"}
        if key:
            env["HTTP_AUTHORIZATION"] = "Bearer " + key
        out = {}
        body = b"".join(server.app(env, lambda s, h: out.update(s=s, h=dict(h))))
        return out["s"], out["h"]["Content-Type"], body

    def test_rest_mcp_and_download(self):
        os.environ["ANVIL_API_KEYS"] = "k1"
        self.addCleanup(os.environ.pop, "ANVIL_API_KEYS", None)
        from anvil import ember
        ember.company_url_candidates = lambda lead, jd="": []
        self.assertTrue(self.call("GET", "/api/v1/stats", key="nope")[0].startswith("401"))
        jd = TestTailorAndBandit.JD
        s, _, b = self.call("POST", "/api/v1/tailor", {"jd_text": jd, "company": "Acme", "role": "Business Analyst"})
        self.assertTrue(s.startswith("200"), b)
        data = json.loads(b)["data"]
        cid = data["canonical_id"]
        pdf_url = data["downloads"]["pdf"].replace("https://x.test", "")
        s, ctype, pdf = self.call("GET", pdf_url)
        self.assertEqual(ctype, "application/pdf")
        self.assertTrue(pdf.startswith(b"%PDF"))
        s, _, b = self.call("POST", "/api/mcp", {"jsonrpc": "2.0", "id": 1, "method": "tools/call",
                                                "params": {"name": "feedback",
                                                           "arguments": {"canonical_id": cid, "outcome": "callback"}}},
                            key=None, qs="key=k1")
        res = json.loads(b)["result"]
        self.assertFalse(res["isError"], res)
        self.assertIn("CALLBACK", res["content"][0]["text"])
        s, _, b = self.call("POST", "/api/mcp", {"jsonrpc": "2.0", "id": 2, "method": "tools/list"})
        self.assertIn("tailor", [t["name"] for t in json.loads(b)["result"]["tools"]])
        self.assertTrue(self.call("GET", "/", key=None)[0].startswith("200"))


if __name__ == "__main__":
    unittest.main()
